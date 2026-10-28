"""Make a masked copy of a real Power BI export for demos, UAT or bug reports.

Names, emails, titles, institutions and URLs are replaced with stable stand-ins, so the same
person still maps to the same stand-in across rows. Manuscript IDs become fake 9-digit IDs.
Counts, dates, stages, tasks and health checks are kept, so the console behaves exactly as it
would on the real file.

    python tests/anonymise.py "path/to/export.xlsx"
    -> writes "path/to/export.anonymised.xlsx"

The output is still a spreadsheet, and .gitignore blocks spreadsheets on purpose. Check the
file by eye before sharing it, and never commit real exports. Needs: pip install openpyxl
"""
import hashlib
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

PERSON = re.compile(r"(full name|reviewer name|referee name|author|task performer name|escalated .*name|paused by full name|assistant)", re.I)
EMAIL = re.compile(r"email", re.I)
FREE_TEXT = re.compile(r"(manuscript title|institution|special issue name|country)", re.I)
URL = re.compile(r"(url|link)", re.I)
MS_ID = re.compile(r"^manuscript id$", re.I)


def token(value, prefix, n=6):
    digest = hashlib.sha256(str(value).strip().lower().encode()).hexdigest()
    return f"{prefix}{digest[:n]}"


def mask(header, value):
    if value in (None, ""):
        return value
    h = str(header or "")
    if MS_ID.search(h):
        return str(int(hashlib.sha256(str(value).encode()).hexdigest(), 16) % 900000000 + 100000000)
    if EMAIL.search(h):
        return f"{token(value, 'user-')}@example.org"
    if PERSON.search(h):
        # keep "A; B; C" author lists as lists
        return "; ".join(f"Person {token(p, '', 5).upper()}" for p in str(value).split(";") if p.strip())
    if FREE_TEXT.search(h):
        return f"[{h} {token(value, '', 5)}]"
    if URL.search(h):
        return "https://review.example.org/details/masked"
    return value


def main(src):
    src = Path(src)
    wb = load_workbook(src)
    for ws in wb.worksheets:
        rows = ws.iter_rows()
        try:
            header = [c.value for c in next(rows)]
        except StopIteration:
            continue
        for row in rows:
            for cell, h in zip(row, header):
                cell.value = mask(h, cell.value)
    out = src.with_name(src.stem + ".anonymised.xlsx")
    wb.save(out)
    print(f"Wrote {out}. Check it before sharing; don't commit it.")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
